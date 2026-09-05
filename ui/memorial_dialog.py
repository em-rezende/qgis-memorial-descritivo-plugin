# -*- coding: utf-8 -*-
import math
import json
import os
import hashlib
import csv
from typing import List, Optional

from ..utils.compat import QtCompat

from qgis.PyQt.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QFormLayout,
    QLineEdit, QTextEdit, QTableWidget, QTableWidgetItem,
    QPushButton, QLabel, QGroupBox, QHeaderView, QSpinBox, QMessageBox,
    QCheckBox, QApplication, QAbstractItemView, QFileDialog,
    QTabWidget, QWidget
)
from qgis.PyQt.QtCore import Qt, QVariant
from qgis.PyQt.QtGui import QColor, QFont
from qgis.gui import QgsRubberBand
from qgis.core import (
    QgsPointXY, QgsGeometry, QgsWkbTypes, QgsProject,
    QgsFeature, QgsField, QgsPalLayerSettings, QgsVectorLayerSimpleLabeling,
    QgsTextFormat, QgsTextBufferSettings
)

from .preview_dialog import PreviewDialog
from .widgets import CheckBoxWidget
from ..core.models import DadosImovel, Segmento, GeometriaImovel
from ..core.geometry import GeometryCalculator
from ..core.database import Database
from ..core.generator import MemorialGenerator
from ..utils.formatters import Formatter


class MemorialDialog(QDialog):
    """Diálogo principal do Memorial Descritivo."""

    # Constantes
    LIMITE_MAX_SEGMENTO_CURVA = 8.0
    TOL_DEFLEXAO_PADRAO = 15

    def __init__(self, iface, feature_id, nodes, crs_name="SIRGAS 2000",
                 crs_auth_id="EPSG:31983", parent=None):
        super().__init__(parent)
        
        self.iface = iface
        self.canvas = iface.mapCanvas()
        self.crs_name = crs_name
        self.crs_auth_id = crs_auth_id
        
        # Dados
        self.raw_nodes = [QgsPointXY(pt.x(), pt.y()) for pt in nodes]
        self.nodes = list(self.raw_nodes)
        self.num_vertices = len(self.raw_nodes)
        self.feature_id = self._gerar_id_unico(feature_id)
        
        # Geometria
        self.geometria = GeometriaImovel()
        self.geometria.num_vertices = self.num_vertices
        self.geometria.is_clockwise = GeometryCalculator.calcular_orientacao(self.raw_nodes)
        self._atualizar_geometria()
        
        # Dados do imóvel
        self.dados = DadosImovel(
            feature_id=self.feature_id,
            raw_nodes=self.raw_nodes
        )
        
        # Banco de dados
        plugin_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        db_path = os.path.join(plugin_dir, "memorial_dados.db")
        self.db = Database(db_path)
        
        # Marcadores visuais no Canvas
        self.rubber_band_vertex = QgsRubberBand(self.canvas, QgsWkbTypes.PointGeometry)
        self.rubber_band_vertex.setColor(QColor(255, 0, 0, 200))
        self.rubber_band_vertex.setIconSize(14)
        self.rubber_band_vertex.setSecondaryStrokeColor(QColor(255, 255, 255, 255))
        self.rubber_band_vertex.setWidth(3)
        self.rubber_band_vertex.reset()
        
        self.rubber_band_segment = QgsRubberBand(self.canvas, QgsWkbTypes.LineGeometry)
        self.rubber_band_segment.setColor(QColor(255, 0, 0, 180))
        self.rubber_band_segment.setWidth(8)
        self.rubber_band_segment.reset()
        
        # Estado interno
        self.segmentos_info = []
        self.confrontantes_salvos = {}
        
        # Inicializa UI e carrega dados
        self._init_ui()
        self._conectar_sinais()
        self._carregar_dados_db()
        self._atualizar_orientacao()
        self._on_vertice_changed(self.spin_v1.value())

    # ------------------------------------------------------------------------
    # Inicialização da Interface
    # ------------------------------------------------------------------------

    def _init_panel_cabecalho(self, parent_layout):
        gb = QGroupBox("Dados do Imóvel / Projeto")
        form = QFormLayout()
        form.setSpacing(6)
     
        # Campos
        self.txt_titulo_proj = QLineEdit("Projeto de Retificação de Área")
        
        # Novos campos LOTE, QUADRA, BAIRRO e MATRÍCULA
        self.txt_lote = QLineEdit("Lote 01")
        self.txt_quadra = QLineEdit("Quadra 01")
        self.txt_bairro = QLineEdit("Bairro Centro")
        
        self.txt_matricula = QLineEdit("00.001")
        self.txt_matricula.setMaximumWidth(100)
     
        self.txt_cidade = QLineEdit("Belo Horizonte")
        self.txt_uf = QLineEdit("MG")
        self.txt_uf.setMaximumWidth(50)
     
        # Linha 1: Lote + Quadra + Bairro + Matrícula
        layout_lote_quadra_bairro = QHBoxLayout()
        layout_lote_quadra_bairro.setContentsMargins(0, 0, 0, 0)
        
        layout_lote_quadra_bairro.addWidget(QLabel("Lote:"))
        layout_lote_quadra_bairro.addWidget(self.txt_lote)
        
        layout_lote_quadra_bairro.addWidget(QLabel("Quadra:"))
        layout_lote_quadra_bairro.addWidget(self.txt_quadra)
        
        layout_lote_quadra_bairro.addWidget(QLabel("Bairro:"))
        layout_lote_quadra_bairro.addWidget(self.txt_bairro)
        
        layout_lote_quadra_bairro.addSpacing(10)
        layout_lote_quadra_bairro.addWidget(QLabel("Matrícula:"))
        layout_lote_quadra_bairro.addWidget(self.txt_matricula)
     
        # Linha 2: Município + UF
        layout_cidade_uf = QHBoxLayout()
        layout_cidade_uf.setContentsMargins(0, 0, 0, 0)
        layout_cidade_uf.addWidget(self.txt_cidade)
        layout_cidade_uf.addWidget(QLabel("UF:"))
        layout_cidade_uf.addWidget(self.txt_uf)
        layout_cidade_uf.addStretch()
     
        # Adiciona ao Form
        form.addRow("Título do Projeto:", self.txt_titulo_proj)
        form.addRow("Localização:", layout_lote_quadra_bairro)
        form.addRow("Município:", layout_cidade_uf)
     
        gb.setLayout(form)
        parent_layout.addWidget(gb)

    def _init_ui(self):
        self.setWindowTitle("Configurar Memorial Descritivo para Registro de Imóveis")
        self.resize(800, 580)
    
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(6)
    
        self.tabs = QTabWidget()
    
        # ==========================================================
        # ABA 1: Dados, Anotações, Assinaturas
        # ==========================================================
        tab1 = QWidget()
        layout_tab1 = QVBoxLayout(tab1)
        layout_tab1.setSpacing(8)
    
        self._init_panel_cabecalho(layout_tab1)      # 1. Dados do imóvel
        self._init_panel_assinaturas(layout_tab1)    # 2. Assinaturas
        self._init_panel_botoes(layout_tab1)         # 3. Preview (fora) + BD (dentro)
        
        layout_tab1.addStretch()
    
        # ==========================================================
        # ABA 2: QUADRO RESUMO
        # ==========================================================
        tab2 = QWidget()
        layout_tab2 = QVBoxLayout(tab2)
        layout_tab2.setSpacing(8)
    
        self._init_panel_vertice(layout_tab2)
        self._init_panel_geometria(layout_tab2)
        self._init_panel_quadro(layout_tab2)
        self._init_panel_desenho(layout_tab2)
    
        # Adiciona as abas ao gerenciador
        self.tabs.addTab(tab1, "📋 Dados, Anotações, Assinaturas")
        self.tabs.addTab(tab2, "📊 Quadro Resumo")
    
        main_layout.addWidget(self.tabs)
    
        # Rodapé Geral (Apenas Sair / Cancelar)
        self._init_rodape(main_layout)

    def _init_panel_vertice(self, parent_layout):
        gb = QGroupBox("Seleção do Vértice Inicial (V1), Sentido Topográfico & Curvatura")
        layout = QHBoxLayout()

        lbl_v = QLabel("V1:")
        self.btn_prev_v = QPushButton("◄ Anterior")
        self.spin_v1 = QSpinBox()
        self.spin_v1.setRange(1, self.num_vertices)
        self.spin_v1.setMinimumWidth(70)
        self.spin_v1.setStyleSheet("font-weight: bold; font-size: 14px;")
        self.btn_next_v = QPushButton("Próximo ►")

        self.btn_inverter_sentido = QPushButton("🔄 Inverter Sentido")
        self.lbl_sentido_status = QLabel("")
        self.lbl_sentido_status.setStyleSheet("font-weight: bold; color: #1565C0;")

        self.chk_curvas = QCheckBox("Detectar Curvas")
        self.chk_curvas.setChecked(True)

        self.lbl_tol = QLabel("Tol. Deflexão (º):")
        self.spin_tol_deflexao = QSpinBox()
        self.spin_tol_deflexao.setRange(1, 45)
        self.spin_tol_deflexao.setValue(self.TOL_DEFLEXAO_PADRAO)
        self.spin_tol_deflexao.setSuffix("°")

        layout.addWidget(lbl_v)
        layout.addWidget(self.btn_prev_v)
        layout.addWidget(self.spin_v1)
        layout.addWidget(self.btn_next_v)
        layout.addSpacing(10)
        layout.addWidget(self.btn_inverter_sentido)
        layout.addWidget(self.lbl_sentido_status)
        layout.addSpacing(15)
        layout.addWidget(self.chk_curvas)
        layout.addWidget(self.lbl_tol)
        layout.addWidget(self.spin_tol_deflexao)
        layout.addStretch()

        gb.setLayout(layout)
        parent_layout.addWidget(gb)

    def _init_panel_geometria(self, parent_layout):
        gb = QGroupBox("Métricas Geométricas do Imóvel")
        gb.setStyleSheet("QGroupBox { font-weight: bold; }")
        layout = QHBoxLayout()

        style_metric = (
            "font-size: 13px; font-weight: bold; color: palette(text); "
            "background-color: palette(base); padding: 6px; border-radius: 4px; "
            "border: 1px solid palette(mid);"
        )

        self.lbl_area_m2 = QLabel("Área: 0,00 m²")
        self.lbl_area_ha = QLabel("0,0000 ha")
        self.lbl_area_alq = QLabel("0,00 alq (SP)")
        self.lbl_perimetro = QLabel("Perímetro: 0,00 m")

        for lbl in [self.lbl_area_m2, self.lbl_area_ha, self.lbl_area_alq, self.lbl_perimetro]:
            lbl.setStyleSheet(style_metric)

        layout.addWidget(self.lbl_area_m2)
        layout.addWidget(self.lbl_area_ha)
        layout.addWidget(self.lbl_area_alq)
        layout.addWidget(self.lbl_perimetro)
        gb.setLayout(layout)
        parent_layout.addWidget(gb)

    
    def _init_panel_quadro(self, parent_layout):
        gb = QGroupBox("Quadro Resumo / Confrontantes por Trecho")
        layout = QVBoxLayout()

        top_layout = QHBoxLayout()
        self.btn_detectar_vizinhos = QPushButton("🔍 Detectar Vizinhos")
        self.btn_exportar_tabela = QPushButton("📊 Exportar Tabela")
        
        self.lbl_crs_info = QLabel(f"<b>Coordenadas:</b> {self.crs_name}")
        self.lbl_crs_info.setStyleSheet("padding-left: 10px; color: palette(text);")
        
        top_layout.addWidget(self.btn_detectar_vizinhos)
        top_layout.addWidget(self.btn_exportar_tabela)
        top_layout.addWidget(self.lbl_crs_info)
        top_layout.addStretch()
        layout.addLayout(top_layout)

        self.table_conf = QTableWidget()
        self.table_conf.setColumnCount(9)
        self.table_conf.horizontalHeader().setVisible(False)
        self.table_conf.verticalHeader().setDefaultSectionSize(20)
        self.table_conf.setStyleSheet("QTableWidget::item { padding: 1px 4px; }")

        try:
            self.table_conf.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
            self.table_conf.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        except AttributeError:
            self.table_conf.setSelectionBehavior(QAbstractItemView.SelectRows)
            self.table_conf.setSelectionMode(QAbstractItemView.SingleSelection)

        layout.addWidget(self.table_conf)
        gb.setLayout(layout)
        parent_layout.addWidget(gb)

    def _init_panel_assinaturas(self, parent_layout):
        gb = QGroupBox("Assinaturas e Detalhes de Confrontantes Gerais")
        form = QFormLayout()
        form.setSpacing(6)
    
        self.txt_proprietario = QTextEdit("Nome do Proprietário / CNPJ 00.000.000/001-00")
        self.txt_proprietario.setMaximumHeight(40) # controla a altura do campo
    
        self.txt_confrontantes_bloco = QTextEdit()
        self.txt_confrontantes_bloco.setMaximumHeight(160) # controla a altura do campo
    
        self.txt_resp_tecnico = QTextEdit("Nome do Responsável Técnico / CAU-BR / CREA / CFT")
        self.txt_resp_tecnico.setMaximumHeight(40) # controla a altura do campo
    
        form.addRow("Proprietário(s):", self.txt_proprietario)
        form.addRow("Bloco de Confrontantes:", self.txt_confrontantes_bloco)
        form.addRow("Responsável Técnico:", self.txt_resp_tecnico)
    
        gb.setLayout(form)
        parent_layout.addWidget(gb)

    def _init_panel_botoes(self, parent_layout):
        """Cria a linha com Preview (fora) e Banco de Dados (dentro da caixa)."""
        
        # Layout horizontal principal
        linha_principal = QHBoxLayout()
        linha_principal.setSpacing(10)
        
        # Botão Pré-visualizar (FORA da caixa, à esquerda)
        self.btn_preview = QPushButton("👁️ Pré-visualizar / Exportar Memorial")
        self.btn_preview.setStyleSheet("font-weight: bold; font-size: 12px;")
        linha_principal.addWidget(self.btn_preview)
        
        # Adiciona um espaçador flexível para empurrar o grupo BD para a direita
        linha_principal.addStretch()
        
        # ========== CAIXA BANCO DE DADOS ==========
        # Cria um GroupBox apenas para os botões do BD
        gb = QGroupBox("Banco de Dados")
        layout_bd = QHBoxLayout()  # Layout horizontal para os botões do BD
        layout_bd.setSpacing(5)
        
        # Botões do Banco de Dados (DENTRO da caixa)
        self.btn_salvar_db_tab1 = QPushButton("💾 Salvar Dados no BD")
        self.btn_limpar_lote_tab1 = QPushButton("🗑️ Limpar Lote Atual")
        self.btn_apagar_tudo_tab1 = QPushButton("⚠️ Apagar Todo o Banco de Dados")
        self.btn_apagar_tudo_tab1.setStyleSheet("color: red; font-weight: bold;")
        
        layout_bd.addWidget(self.btn_salvar_db_tab1)
        layout_bd.addWidget(self.btn_limpar_lote_tab1)
        layout_bd.addWidget(self.btn_apagar_tudo_tab1)
        
        gb.setLayout(layout_bd)
        linha_principal.addWidget(gb)  # Adiciona a caixa BD à linha principal
        
        # Adiciona a linha principal ao layout pai
        parent_layout.addLayout(linha_principal)

    def _init_panel_desenho(self, parent_layout):
        h_layout = QHBoxLayout()
        h_layout.setSpacing(10)
        
        self._criar_grupo_anotacao(h_layout)
        self._criar_grupo_bd(h_layout)
        
        parent_layout.addLayout(h_layout)
    
    def _criar_grupo_anotacao(self, parent_layout):
        gb = QGroupBox("Ferramentas de Anotação no Mapa (QGIS Canvas)")
        layout = QHBoxLayout()
        layout.setSpacing(5)
        
        self.btn_acrescentar_desenho = QPushButton("Acrescentar no Desenho")
        self.btn_atualizar_desenho = QPushButton("Atualizar")
        self.btn_remover_desenho = QPushButton("Remover do Desenho")
        self.btn_remover_desenho.setStyleSheet("color: red;")
        
        layout.addWidget(self.btn_acrescentar_desenho)
        layout.addWidget(self.btn_atualizar_desenho)
        layout.addWidget(self.btn_remover_desenho)
        layout.addStretch()
        gb.setLayout(layout)
        parent_layout.addWidget(gb, 1)

    def _criar_grupo_bd(self, parent_layout):
        """Cria o grupo BANCO DE DADOS para a Aba 2 (Quadro Resumo)."""
        gb = QGroupBox("Banco de Dados")
        layout = QHBoxLayout()
        layout.setSpacing(5)
        
        self.btn_salvar_db = QPushButton("💾 Salvar Dados no BD")
        self.btn_limpar_lote = QPushButton("🗑️ Limpar Lote Atual")
        self.btn_apagar_tudo = QPushButton("⚠️ Apagar Todo o Banco de Dados")
        self.btn_apagar_tudo.setStyleSheet("color: red; font-weight: bold;")
        
        layout.addWidget(self.btn_salvar_db)
        layout.addWidget(self.btn_limpar_lote)
        layout.addWidget(self.btn_apagar_tudo)
        layout.addStretch()
        gb.setLayout(layout)
        parent_layout.addWidget(gb, 1)

   
    def _init_rodape(self, parent_layout):
        layout = QHBoxLayout()
        layout.addStretch()

        self.btn_cancelar = QPushButton("Sair / Cancelar")
        self.btn_cancelar.setStyleSheet("font-weight: bold; min-width: 100px;")

        layout.addWidget(self.btn_cancelar)
        parent_layout.addLayout(layout)

    def _conectar_sinais(self):
        self.spin_v1.valueChanged.connect(self._on_vertice_changed)
        self.btn_prev_v.clicked.connect(self._vertice_anterior)
        self.btn_next_v.clicked.connect(self._proximo_vertice)
        self.btn_inverter_sentido.clicked.connect(self._inverter_sentido)
    
        self.chk_curvas.toggled.connect(self._on_curvas_toggled)
        self.spin_tol_deflexao.valueChanged.connect(
            lambda: self._on_vertice_changed(self.spin_v1.value())
        )
    
        self.btn_detectar_vizinhos.clicked.connect(self._detectar_confrontantes)
        
        # Sinais dos botões do Banco de Dados (Aba 2)
        self.btn_salvar_db.clicked.connect(self._salvar_dados_db)
        self.btn_limpar_lote.clicked.connect(self._limpar_lote_atual)
        self.btn_apagar_tudo.clicked.connect(self._apagar_tudo)
    
        # Sinais dos botões do Banco de Dados (Aba 1) - ADICIONE ESTAS LINHAS
        self.btn_salvar_db_tab1.clicked.connect(self._salvar_dados_db)
        self.btn_limpar_lote_tab1.clicked.connect(self._limpar_lote_atual)
        self.btn_apagar_tudo_tab1.clicked.connect(self._apagar_tudo)
    
        self.btn_acrescentar_desenho.clicked.connect(self._acrescentar_anotacoes)
        self.btn_atualizar_desenho.clicked.connect(self._atualizar_anotacoes)
        self.btn_remover_desenho.clicked.connect(self._remover_anotacoes)
    
        self.btn_preview.clicked.connect(self._exibir_preview)
        self.btn_exportar_tabela.clicked.connect(self._exportar_tabela)
        self.btn_cancelar.clicked.connect(self.reject)
    
        self.table_conf.itemChanged.connect(self._ao_alterar_item_tabela)
        self.table_conf.itemSelectionChanged.connect(self._destacar_segmento)
        self.table_conf.cellClicked.connect(lambda r, c: self._destacar_segmento())
        self.table_conf.currentCellChanged.connect(
            lambda r, c, pr, pc: self._destacar_segmento()
        )

    # ------------------------------------------------------------------------
    # Métodos Principais
    # ------------------------------------------------------------------------

    def _get_lote_quadra_concatenado(self) -> str:
        """Helper para obter a string concatenada de Lote, Quadra e Bairro."""
        p_lote = self.txt_lote.text().strip()
        p_quadra = self.txt_quadra.text().strip()
        p_bairro = self.txt_bairro.text().strip()
        
        partes = [p for p in [p_lote, p_quadra, p_bairro] if p]
        return " - ".join(partes)

    def _gerar_id_unico(self, feature_id):
        str_id = str(feature_id).strip() if feature_id is not None else ""
        if str_id and str_id not in ["None", "0", "-1", "NULL"]:
            return f"FEAT_{str_id}"
        coord_str = "_".join([f"{p.x():.3f},{p.y():.3f}" for p in self.raw_nodes])
        return f"GEO_{hashlib.md5(coord_str.encode('utf-8')).hexdigest()[:12]}"

    def _atualizar_geometria(self):
        self.geometria.area_m2, self.geometria.perimetro = GeometryCalculator.calcular_area_perimetro(self.nodes)
        self.geometria.is_clockwise = GeometryCalculator.calcular_orientacao(self.nodes)

    def _atualizar_metricas(self):
        dados_formatados = Formatter.formatar_area(self.geometria.area_m2)
        self.lbl_area_m2.setText(f"Área: {dados_formatados['m2']} m²")
        self.lbl_area_ha.setText(f"{dados_formatados['ha']} ha")
        self.lbl_area_alq.setText(f"{dados_formatados['alq_sp']} alq (SP)")
        self.lbl_perimetro.setText(
            f"Perímetro: {Formatter.formatar_numero_br(self.geometria.perimetro, 2)} m"
        )

    def _atualizar_orientacao(self):
        self.nodes = GeometryCalculator.reorientar_nodes(self.raw_nodes, self.geometria.is_clockwise)
        self._atualizar_geometria()
        self.lbl_sentido_status.setText(
            "Sentido: HORÁRIO" if self.geometria.is_clockwise else "Sentido: ANTI-HORÁRIO"
        )

    def _on_vertice_changed(self, v1_val):
        self._salvar_estado_tabela()
    
        idx_start = v1_val - 1
        pt_v1 = self.nodes[idx_start]
        self.rubber_band_vertex.setToGeometry(QgsGeometry.fromPointXY(pt_v1), None)
        self.canvas.refresh()
    
        nodes_ord = self.nodes[idx_start:] + self.nodes[:idx_start]
        self.nodes_ordenados = nodes_ord
    
        self.geometria.area_m2, self.geometria.perimetro = GeometryCalculator.calcular_area_perimetro(nodes_ord)
        self._atualizar_metricas()
    
        curvas_ativas = self.chk_curvas.isChecked()

        if curvas_ativas:
            segs_raw = GeometryCalculator.processar_segmentos(
                nodes_ord,
                True,
                float(self.spin_tol_deflexao.value())
            )
        else:
            segs_raw = []
            num_pts = len(nodes_ord)
            for i in range(num_pts):
                p_atual = nodes_ord[i]
                p_prox = nodes_ord[(i + 1) % num_pts]
                
                v_de = f"V{i + 1}"
                v_para = f"V{1 if i == num_pts - 1 else i + 2}"
                
                dx = p_prox.x() - p_atual.x()
                dy = p_prox.y() - p_atual.y()
                dist = math.hypot(dx, dy)
                
                az_rad = math.atan2(dx, dy)
                if az_rad < 0:
                    az_rad += 2 * math.pi
                
                graus_tot = math.degrees(az_rad)
                g = int(graus_tot)
                m = int((graus_tot - g) * 60)
                s = (graus_tot - g - m / 60.0) * 3600.0
                az_str = f'{g:02d}°{m:02d}\'{s:05.2f}"'
                
                segs_raw.append({
                    'de': v_de,
                    'para': v_para,
                    'p1': p_atual,
                    'p2': p_prox,
                    'dist': dist,
                    'az': az_str,
                    'is_curva': False,
                    'raio': 0.0,
                    'arco': 0.0,
                    'corda': 0.0,
                    'sentido_curva': ""
                })
    
        self.segmentos_info = []
    
        for seg in segs_raw:
            seg_key = GeometryCalculator.gerar_chave_segmento(seg['p1'], seg['p2'])
            v_de = seg['de']
            v_para = seg['para']
            conf_padrao = f"CONFRONTANTE TRECHO {v_de}-{v_para}"
            requer_ass = True
    
            if seg_key in self.confrontantes_salvos:
                dados_s = self.confrontantes_salvos[seg_key]
                conf_padrao = dados_s.get('nome', conf_padrao)
                requer_ass = dados_s.get('assinatura', True)
    
            az_value = seg.get('az', '00º00\'00.00"')
    
            self.segmentos_info.append(Segmento(
                seg_key=seg_key,
                de=v_de,
                para=v_para,
                p1_n=seg['p1'].y(),
                p1_e=seg['p1'].x(),
                p2_n=seg['p2'].y(),
                p2_e=seg['p2'].x(),
                az=az_value,
                dist=seg['dist'],
                is_curva=seg.get('is_curva', False),
                raio=seg.get('raio', 0.0),
                arco=seg.get('arco', 0.0),
                corda=seg.get('corda', 0.0),
                sentido_curva=seg.get('sentido_curva', ""),
                confrontante=conf_padrao,
                requer_assinatura=requer_ass
            ))
    
        self._reconstruir_tabela()

    def _salvar_estado_tabela(self):
        if not hasattr(self, 'segmentos_info') or not self.segmentos_info:
            return

        for idx, seg in enumerate(self.segmentos_info):
            row = idx + 1
            item_conf = self.table_conf.item(row, 8)
            chk_widget = self.table_conf.cellWidget(row, 7)

            requer_ass = True
            if chk_widget:
                if hasattr(chk_widget, 'is_checked'):
                    requer_ass = chk_widget.is_checked()
                else:
                    chk = chk_widget.findChild(QCheckBox)
                    if chk:
                        requer_ass = chk.isChecked()

            nome_conf = item_conf.text() if item_conf else f"CONFRONTANTE TRECHO {seg.de}-{seg.para}"

            self.confrontantes_salvos[seg.seg_key] = {
                "nome": nome_conf,
                "assinatura": requer_ass
            }

    def _reconstruir_tabela(self):
        self.table_conf.blockSignals(True)
        
        flags_read_only = QtCompat.get_item_flags_read_only()
        flag_edit = QtCompat.get_item_flags_editable()
        align_center = QtCompat.get_align_center()
        
        headers = [
            ("VÉRTICE", 0),
            ("NORTE", 1),
            ("ESTE", 2),
            ("DE", 3),
            ("PARA", 4),
            ("AZIMUTE / TIPO", 5),
            ("DISTÂNCIA (m)", 6),
            ("ASSINATURA", 7),
            ("CONFRONTANTE", 8)
        ]
        
        self.table_conf.setRowCount(1 + len(self.segmentos_info))
        
        for text, col in headers:
            it = QTableWidgetItem(text)
            it.setFlags(flags_read_only)
            it.setTextAlignment(align_center)
            self.table_conf.setItem(0, col, it)
        
        for idx, seg in enumerate(self.segmentos_info):
            row = idx + 1

            if seg.is_curva:
                az_display = f"CURVA - R={Formatter.formatar_numero_br(seg.raio, 2)}m"
            else:
                az_display = seg.az

            items = [
                (seg.de, flags_read_only),
                (seg.norte_str, flags_read_only),
                (seg.este_str, flags_read_only),
                (seg.de, flags_read_only),
                (seg.para, flags_read_only),
                (az_display, flags_read_only),
                (Formatter.formatar_numero_br(seg.dist, 2), flags_read_only),
            ]

            for col_idx, (val, flags) in enumerate(items):
                it = QTableWidgetItem(val)
                it.setFlags(flags)
                it.setTextAlignment(align_center)
                self.table_conf.setItem(row, col_idx, it)

            chk_widget = CheckBoxWidget(seg.requer_assinatura)
            self.table_conf.setCellWidget(row, 7, chk_widget)

            it_conf = QTableWidgetItem(seg.confrontante)
            it_conf.setFlags(flag_edit)
            self.table_conf.setItem(row, 8, it_conf)

        header = self.table_conf.horizontalHeader()
        for col in range(8):
            try:
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.ResizeToContents)
            except AttributeError:
                header.setSectionResizeMode(col, QHeaderView.ResizeToContents)

        try:
            header.setSectionResizeMode(8, QHeaderView.ResizeMode.Stretch)
        except AttributeError:
            header.setSectionResizeMode(8, QHeaderView.Stretch)

        self.table_conf.blockSignals(False)
        self._atualizar_bloco_confrontantes()

    def _atualizar_bloco_confrontantes(self):
        confrontantes_unicos = []

        for row in range(1, self.table_conf.rowCount()):
            chk_widget = self.table_conf.cellWidget(row, 7)
            requer_assinatura = True
            if chk_widget:
                requer_assinatura = chk_widget.is_checked()

            item_conf = self.table_conf.item(row, 8)
            if item_conf and requer_assinatura:
                texto = item_conf.text().strip()
                if texto and texto not in confrontantes_unicos:
                    confrontantes_unicos.append(texto)

        bloco_txt = "\n\n".join([f"__________________________________________________\n{conf}" for conf in confrontantes_unicos])
        self.txt_confrontantes_bloco.blockSignals(True)
        self.txt_confrontantes_bloco.setPlainText(bloco_txt)
        self.txt_confrontantes_bloco.blockSignals(False)

    # ------------------------------------------------------------------------
    # Eventos da UI
    # ------------------------------------------------------------------------

    def _vertice_anterior(self):
        self._salvar_estado_tabela()
        val = self.spin_v1.value() - 1
        if val < 1:
            val = self.num_vertices
        self.spin_v1.setValue(val)

    def _proximo_vertice(self):
        self._salvar_estado_tabela()
        val = self.spin_v1.value() + 1
        if val > self.num_vertices:
            val = 1
        self.spin_v1.setValue(val)

    def _inverter_sentido(self):
        self._salvar_estado_tabela()
        self.geometria.is_clockwise = not self.geometria.is_clockwise
        self._atualizar_orientacao()
        self._atualizar_metricas()
        self._on_vertice_changed(self.spin_v1.value())

    def _on_curvas_toggled(self, checked):
        self.spin_tol_deflexao.setEnabled(checked)
        self._on_vertice_changed(self.spin_v1.value())

    def _ao_alterar_item_tabela(self, item):
        if item.column() == 8:
            self._salvar_estado_tabela()
            self._atualizar_bloco_confrontantes()

    def _destacar_segmento(self):
        row_atual = self.table_conf.currentRow()
        if row_atual < 1:
            self.rubber_band_segment.reset()
            self.canvas.refresh()
            return

        seg_idx = row_atual - 1
        if 0 <= seg_idx < len(self.segmentos_info):
            seg = self.segmentos_info[seg_idx]
            p1 = QgsPointXY(seg.p1_e, seg.p1_n)
            p2 = QgsPointXY(seg.p2_e, seg.p2_n)
            geom = QgsGeometry.fromPolylineXY([p1, p2])
            self.rubber_band_segment.setToGeometry(geom, None)
            self.canvas.refresh()

    # ------------------------------------------------------------------------
    # Banco de Dados
    # ------------------------------------------------------------------------

    def _carregar_dados_db(self):
        dados = self.db.carregar(self.feature_id)
        if dados:
            self.spin_v1.setValue(dados["v1_index"])
            self.txt_titulo_proj.setText(dados["titulo_proj"])
            
            # Se existirem os novos campos gravados, preenche diretamente
            if dados.get("lote") or dados.get("quadra") or dados.get("bairro"):
                self.txt_lote.setText(dados.get("lote", ""))
                self.txt_quadra.setText(dados.get("quadra", ""))
                self.txt_bairro.setText(dados.get("bairro", ""))
            else:
                # Fallback para string antiga salva em lote_quadra
                self.txt_lote.setText(dados.get("lote_quadra", ""))
                
            self.txt_matricula.setText(dados["matricula"])
            self.txt_cidade.setText(dados["cidade"])
            self.txt_uf.setText(dados["uf"])
            self.txt_proprietario.setPlainText(dados["proprietario"])
            self.txt_resp_tecnico.setPlainText(dados["resp_tecnico"])
            self.confrontantes_salvos = dados["confrontantes"]

    def _salvar_dados_db(self, silencioso=False):
        try:
            self._salvar_estado_tabela()
            lote_quadra_conc = self._get_lote_quadra_concatenado()
        
            self.db.salvar(
                feature_id=self.feature_id,
                v1_index=self.spin_v1.value(),
                titulo_proj=self.txt_titulo_proj.text(),
                lote=self.txt_lote.text(),
                quadra=self.txt_quadra.text(),
                bairro=self.txt_bairro.text(),
                lote_quadra=lote_quadra_conc,
                matricula=self.txt_matricula.text(),
                cidade=self.txt_cidade.text(),
                uf=self.txt_uf.text(),
                proprietario=self.txt_proprietario.toPlainText(),
                resp_tecnico=self.txt_resp_tecnico.toPlainText(),
                confrontantes_salvos=self.confrontantes_salvos,
                raw_nodes=self.raw_nodes
            )
        
            if not silencioso:
                QMessageBox.information(self, "Banco de Dados", "Dados salvos com sucesso!")
        except Exception as e:
            QMessageBox.critical(self, "Erro no Banco de Dados", f"Erro ao salvar:\n{str(e)}")

    def _limpar_lote_atual(self):
        resposta = QMessageBox.question(
            self,
            "Confirmar Exclusão",
            "Deseja realmente apagar os dados salvos do lote atual?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if resposta == QMessageBox.Yes:
            self.db.excluir(self.feature_id)
            self.confrontantes_salvos = {}
            self._on_vertice_changed(self.spin_v1.value())
            QMessageBox.information(self, "Banco de Dados", "Dados removidos com sucesso!")

    def _apagar_tudo(self):
        resposta = QMessageBox.warning(
            self,
            "⚠️ ATENÇÃO: Apagar Todo o Banco de Dados",
            "Esta ação excluirá PERMANENTEMENTE as informações de TODOS os lotes.\n\nTem certeza?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if resposta == QMessageBox.Yes:
            self.db.excluir_todos()
            self.confrontantes_salvos = {}
            self._on_vertice_changed(self.spin_v1.value())
            QMessageBox.information(self, "Banco de Dados", "Banco de dados completamente apagado!")

    # ------------------------------------------------------------------------
    # Detectar Vizinhos
    # ------------------------------------------------------------------------

    def _detectar_confrontantes(self):
        outros_lotes = self.db.listar_outros_lotes(self.feature_id)

        if not outros_lotes:
            QMessageBox.information(self, "Detectar Vizinhos", "Nenhum outro lote cadastrado.")
            return

        encontrados = 0

        for row_idx, seg in enumerate(self.segmentos_info):
            p1 = QgsPointXY(seg.p1_e, seg.p1_n)
            p2 = QgsPointXY(seg.p2_e, seg.p2_n)
            line_seg = QgsGeometry.fromPolylineXY([p1, p2])
            dist_seg = line_seg.length()

            if dist_seg > 0.40:
                pts_core = []
                num_samples = 10
                for step in range(num_samples + 1):
                    d = 0.20 + (dist_seg - 0.40) * (step / num_samples)
                    pt = line_seg.interpolate(d).asPoint()
                    pts_core.append(pt)
                line_core = QgsGeometry.fromPolylineXY(pts_core)
            else:
                line_core = line_seg

            buffer_core = line_core.buffer(0.02, 3)

            for f_id, lote_q, mat, prop, nodes_json_str in outros_lotes:
                if not nodes_json_str:
                    continue
                try:
                    pts_dict = json.loads(nodes_json_str)
                    pts_vizinho = [QgsPointXY(p['x'], p['y']) for p in pts_dict]
                    geom_vizinho = QgsGeometry.fromPolylineXY(pts_vizinho + [pts_vizinho[0]])
                    intersecao = geom_vizinho.intersection(buffer_core)

                    if not intersecao.isEmpty() and intersecao.length() > 0.20:
                        partes = []
                        if lote_q and lote_q.strip():
                            partes.append(lote_q.strip().upper())
                        if mat and mat.strip():
                            partes.append(f"MATRÍCULA: {mat.strip()}")
                        if prop and prop.strip():
                            partes.append(f"PROP: {prop.strip().upper()}")

                        novo_conf = " - ".join(partes) if partes else f"LOTE ID: {f_id}"

                        item = self.table_conf.item(row_idx + 1, 8)
                        if item:
                            item.setText(novo_conf)
                            encontrados += 1
                        break
                except Exception:
                    continue

        self._salvar_estado_tabela()
        if encontrados > 0:
            QMessageBox.information(self, "Detectar Vizinhos", f"{encontrados} confrontante(s) detectado(s).")
        else:
            QMessageBox.information(self, "Detectar Vizinhos", "Nenhuma divisa coincidente identificada.")

    # ------------------------------------------------------------------------
    # Anotações no Mapa
    # ------------------------------------------------------------------------

    def _acrescentar_anotacoes(self):
        try:
            from ..core.annotator import MapAnnotator
        except ImportError:
            from core.annotator import MapAnnotator
        
        dados = self._get_dados()
        
        if not dados['segmentos']:
            QMessageBox.warning(self, "Anotações", "Não há segmentos para anotar.")
            return

        id_imovel = (dados['lote_quadra'].strip() or dados['matricula'].strip())
    
        annotator = MapAnnotator(self.crs_auth_id)
        
        try:
            if id_imovel:
                annotator.remover_anotacoes(id_imovel=id_imovel)
            else:
                annotator.remover_anotacoes()
        except Exception:
            pass
    
        # Adiciona cada vértice apenas uma vez (do ponto inicial de cada segmento V1, V2, ..., VN)
        vertices = []
        for seg in dados['segmentos']:
            vertices.append({"nome": seg.de, "x": seg.p1_e, "y": seg.p1_n})

        # A instrução que adicionava o "V1" duplicado ao final foi removida daqui.
    
        try:
            annotator.criar_camada_pontos(vertices, id_imovel=id_imovel)
            annotator.criar_camada_segmentos(dados['segmentos'], id_imovel=id_imovel)
            annotator.criar_camada_centro(dados, self.geometria, id_imovel=id_imovel)
            
            self.canvas.refresh()
            QMessageBox.information(self, "Anotações", "Anotações adicionadas ao Canvas!")
        except Exception as e:
            QMessageBox.critical(self, "Erro", f"Erro ao criar anotações:\n{str(e)}")

    def _atualizar_anotacoes(self):
        self._acrescentar_anotacoes()

    def _remover_anotacoes(self):
        try:
            from ..core.annotator import MapAnnotator
        except ImportError:
            from core.annotator import MapAnnotator
            
        try:
            dados = self._get_dados()
            id_imovel = (dados['lote_quadra'].strip() or dados['matricula'].strip())

            annotator = MapAnnotator(self.crs_auth_id)
            
            if id_imovel:
                annotator.remover_anotacoes(id_imovel=id_imovel)
            else:
                annotator.remover_anotacoes()

            self.canvas.refresh()
            QMessageBox.information(self, "Anotações", "Anotações do imóvel removidas do Canvas!")
        except Exception as e:
            QMessageBox.warning(self, "Anotações", f"Não foi possível remover as camadas:\n{str(e)}")

    # ------------------------------------------------------------------------
    # Geração e Exportação
    # ------------------------------------------------------------------------

    def _get_dados(self) -> dict:
        self._salvar_estado_tabela()

        segmentos_atualizados = []
        for idx, seg in enumerate(self.segmentos_info):
            item_conf = self.table_conf.item(idx + 1, 8)
            chk_widget = self.table_conf.cellWidget(idx + 1, 7)

            requer_ass = True
            if chk_widget:
                requer_ass = chk_widget.is_checked()

            seg_copy = Segmento(
                seg_key=seg.seg_key,
                de=seg.de,
                para=seg.para,
                p1_n=seg.p1_n,
                p1_e=seg.p1_e,
                p2_n=seg.p2_n,
                p2_e=seg.p2_e,
                az=seg.az,
                dist=seg.dist,
                is_curva=seg.is_curva,
                raio=seg.raio,
                arco=seg.arco,
                corda=seg.corda,
                sentido_curva=seg.sentido_curva,
                confrontante=item_conf.text() if item_conf else seg.confrontante,
                requer_assinatura=requer_ass
            )
            segmentos_atualizados.append(seg_copy)

        return {
            "titulo_proj": self.txt_titulo_proj.text(),
            "lote": self.txt_lote.text(),
            "quadra": self.txt_quadra.text(),
            "bairro": self.txt_bairro.text(),
            "lote_quadra": self._get_lote_quadra_concatenado(),
            "matricula": self.txt_matricula.text(),
            "cidade": self.txt_cidade.text(),
            "uf": self.txt_uf.text(),
            "proprietario": self.txt_proprietario.toPlainText(),
            "confrontantes_bloco": self.txt_confrontantes_bloco.toPlainText(),
            "resp_tecnico": self.txt_resp_tecnico.toPlainText(),
            "segmentos": segmentos_atualizados
        }

    def _gerar_texto_memorial(self, formato="txt"):
        dados = self._get_dados()

        self._atualizar_bloco_confrontantes()
        dados["confrontantes_bloco"] = self.txt_confrontantes_bloco.toPlainText()

        dados_header = {
            "titulo_proj": dados["titulo_proj"],
            "lote": dados["lote"],
            "quadra": dados["quadra"],
            "bairro": dados["bairro"],
            "lote_quadra": dados["lote_quadra"],
            "matricula": dados["matricula"],
            "cidade": dados["cidade"],
            "uf": dados["uf"],
            "proprietario": dados["proprietario"],
            "confrontantes_bloco": dados["confrontantes_bloco"],
            "resp_tecnico": dados["resp_tecnico"]
        }

        segmentos_dict = []
        for s in dados["segmentos"]:
            segmentos_dict.append({
                "de": s.de,
                "para": s.para,
                "p1_n": s.p1_n,
                "p1_e": s.p1_e,
                "p2_n": s.p2_n,
                "p2_e": s.p2_e,
                "is_curva": s.is_curva,
                "sentido_curva": s.sentido_curva,
                "raio": s.raio,
                "arco": s.arco,
                "corda": s.corda,
                "az": s.az,
                "dist": s.dist,
                "confrontante": s.confrontante,
                "requer_assinatura": s.requer_assinatura
            })

        return MemorialGenerator.gerar_texto(
            dados_header,
            segmentos_dict,
            self.geometria.area_m2,
            self.geometria.perimetro
        )

    def _exibir_preview(self):
        texto = self._gerar_texto_memorial()
        dlg = PreviewDialog(texto, self)
        dlg.exec()

    def _exportar_tabela(self):
        dados = self._get_dados()
        titulo = dados['lote_quadra'].strip().upper()
        area_texto = f"ÁREA = {Formatter.formatar_numero_br(self.geometria.area_m2, 2)}m²"

        caminho, _ = QFileDialog.getSaveFileName(
            self,
            "Exportar Tabela Resumo",
            f"Quadro_Resumo_{titulo.replace(' ', '_')}",
            "Arquivo CSV (*.csv);;Planilha Excel (*.xlsx)"
        )

        if not caminho:
            return

        linhas = [
            [titulo, "", "", "", "", "", ""],
            ["QUADRO RESUMO", "", "", "", "", "", ""],
            ["VÉRTICE", f"COORDENADAS ({self.crs_name})", "", "DE", "PARA", "AZIMUTE/RAIO", "DISTÂNCIA (m)"],
            ["", "NORTE", "ESTE", "", "", "", ""]
        ]

        for seg in dados['segmentos']:
            az_raio = Formatter.formatar_numero_br(seg.raio, 2) if seg.is_curva else seg.az
            linhas.append([
                seg.de,
                seg.norte_str,
                seg.este_str,
                seg.de,
                seg.para,
                az_raio,
                Formatter.formatar_numero_br(seg.dist, 2)
            ])

        linhas.append([area_texto, "", "", "", "", "", ""])

        if caminho.lower().endswith(".xlsx"):
            self._exportar_xlsx(caminho, linhas)
        else:
            self._exportar_csv(caminho, linhas)

    def _exportar_csv(self, caminho, linhas):
        try:
            with open(caminho, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f, delimiter=';')
                writer.writerows(linhas)
            QMessageBox.information(self, "Exportar", "CSV exportado com sucesso!")
        except Exception as e:
            QMessageBox.critical(self, "Erro", f"Erro ao salvar: {str(e)}")

    def _exportar_xlsx(self, caminho, linhas):
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, Border, Side
        except ImportError:
            QMessageBox.warning(self, "Biblioteca Ausente", "Instale 'openpyxl' para exportar XLSX.")
            self._exportar_csv(caminho.rsplit('.', 1)[0] + ".csv", linhas)
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Quadro Resumo"

            for linha in linhas:
                ws.append(linha)

            font_bold = Font(name="Arial", size=10, bold=True)
            font_normal = Font(name="Arial", size=10)
            align_center = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
                for cell in row:
                    cell.font = font_normal
                    cell.alignment = align_center
                    cell.border = thin_border

            ws.merge_cells('A1:G1')
            ws.merge_cells('A2:G2')
            ws.merge_cells('B3:C3')
            ws.merge_cells('A3:A4')
            ws.merge_cells('D3:D4')
            ws.merge_cells('E3:E4')
            ws.merge_cells('F3:F4')
            ws.merge_cells('G3:G4')

            last_row = ws.max_row
            ws.merge_cells(start_row=last_row, start_column=1, end_row=last_row, end_column=7)
            ws.cell(row=last_row, column=1).font = font_bold

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(caminho)
            QMessageBox.information(self, "Exportar", "XLSX exportado com sucesso!")
        except Exception as e:
            QMessageBox.critical(self, "Erro", f"Erro ao criar XLSX: {str(e)}")

    # ------------------------------------------------------------------------
    # Eventos de Fechamento
    # ------------------------------------------------------------------------

    def accept(self):
        self._salvar_dados_db(silencioso=True)
        self.rubber_band_vertex.reset()
        self.rubber_band_segment.reset()
        self.canvas.refresh()
        super().accept()

    def closeEvent(self, event):
        self.rubber_band_vertex.reset()
        self.rubber_band_segment.reset()
        self.canvas.refresh()
        super().closeEvent(event)
    
    def reject(self):
        self.rubber_band_vertex.reset()
        self.rubber_band_segment.reset()
        self.canvas.refresh()
        super().reject()
        